from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

import main


VALID_CSV = """#Data operacji;#Opis operacji;#Rachunek;#Kategoria;#Kwota
2026-01-05;Shop   Card payment;123;Food;-12,34 PLN
2026-01-06;Refund;123;Food;5,00 PLN
"""

INVALID_CSV = """not an mbank export
1;2;3
"""


class FakeRoot:
    def withdraw(self) -> None:
        pass

    def update(self) -> None:
        pass

    def destroy(self) -> None:
        pass


class BatchParsingTest(unittest.TestCase):
    def test_select_input_files_returns_all_selected_paths(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            first_file = Path(temp_dir) / "first.csv"
            second_file = Path(temp_dir) / "second.csv"

            with (
                mock.patch.object(main, "Tk", return_value=FakeRoot()),
                mock.patch.object(
                    main.filedialog,
                    "askopenfilenames",
                    return_value=(str(first_file), str(second_file)),
                ),
            ):
                self.assertEqual(
                    main.select_input_files(),
                    [first_file, second_file],
                )

    def test_parse_input_files_continues_after_individual_failures(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            valid_first = self.write_file(temp_path / "first.csv", VALID_CSV)
            invalid_second = self.write_file(temp_path / "second.csv", INVALID_CSV)
            valid_third = self.write_file(temp_path / "third.csv", VALID_CSV)
            invalid_fourth = self.write_file(temp_path / "fourth.csv", INVALID_CSV)
            valid_fifth = self.write_file(temp_path / "fifth.csv", VALID_CSV)

            parsed_files, failed_files = main.parse_input_files(
                [
                    valid_first,
                    invalid_second,
                    valid_third,
                    invalid_fourth,
                    valid_fifth,
                ]
            )

            self.assertEqual(
                [parsed_file.input_file for parsed_file in parsed_files],
                [valid_first, valid_third, valid_fifth],
            )
            self.assertEqual(
                [failed_file.input_file for failed_file in failed_files],
                [invalid_second, invalid_fourth],
            )
            self.assertTrue(main.output_file_for(valid_first).exists())
            self.assertTrue(main.output_file_for(valid_third).exists())
            self.assertTrue(main.output_file_for(valid_fifth).exists())
            self.assertFalse(main.output_file_for(invalid_second).exists())
            self.assertFalse(main.output_file_for(invalid_fourth).exists())
            self.assertTrue(
                all(
                    isinstance(failed_file.error, main.CsvFormatError)
                    for failed_file in failed_files
                )
            )

            workbook = load_workbook(main.output_file_for(valid_first))
            self.assertEqual(workbook.sheetnames, ["Styczeń 2026"])

    def test_batch_summary_detail_lists_outputs_and_failures(self) -> None:
        output_file = Path("/tmp/first_parsed.xlsx")
        failed_file = Path("/tmp/second.csv")

        detail = main.batch_summary_detail(
            [
                main.ParsedFile(
                    input_file=Path("/tmp/first.csv"),
                    output_file=output_file,
                    row_count=3,
                )
            ],
            [
                main.FailedFile(
                    input_file=failed_file,
                    error=main.CsvFormatError("missing header"),
                )
            ],
        )

        self.assertIn("Created outputs:", detail)
        self.assertIn(str(output_file), detail)
        self.assertIn("Failed files:", detail)
        self.assertIn(str(failed_file), detail)
        self.assertIn("missing header", detail)

    def test_show_batch_summary_dialog_uses_info_when_all_files_parse(self) -> None:
        parsed_file = main.ParsedFile(
            input_file=Path("/tmp/first.csv"),
            output_file=Path("/tmp/first_parsed.xlsx"),
            row_count=2,
        )

        with (
            mock.patch.object(main, "Tk", return_value=FakeRoot()),
            mock.patch.object(main.messagebox, "showinfo") as showinfo,
            mock.patch.object(main.messagebox, "showwarning") as showwarning,
            mock.patch.object(main.messagebox, "showerror") as showerror,
        ):
            main.show_batch_summary_dialog([parsed_file], [])

        showinfo.assert_called_once()
        self.assertEqual(showinfo.call_args.kwargs["title"], "Parsing complete")
        self.assertEqual(showinfo.call_args.kwargs["message"], "Parsed 1 of 1 file.")
        showwarning.assert_not_called()
        showerror.assert_not_called()

    def test_show_batch_summary_dialog_uses_warning_for_partial_failures(self) -> None:
        parsed_file = main.ParsedFile(
            input_file=Path("/tmp/first.csv"),
            output_file=Path("/tmp/first_parsed.xlsx"),
            row_count=2,
        )
        failed_file = main.FailedFile(
            input_file=Path("/tmp/second.csv"),
            error=main.CsvFormatError("missing header"),
        )

        with (
            mock.patch.object(main, "Tk", return_value=FakeRoot()),
            mock.patch.object(main.messagebox, "showinfo") as showinfo,
            mock.patch.object(main.messagebox, "showwarning") as showwarning,
            mock.patch.object(main.messagebox, "showerror") as showerror,
        ):
            main.show_batch_summary_dialog([parsed_file], [failed_file])

        showwarning.assert_called_once()
        self.assertEqual(
            showwarning.call_args.kwargs["title"],
            "Parsing completed with errors",
        )
        self.assertEqual(showwarning.call_args.kwargs["message"], "Parsed 1 of 2 files.")
        showinfo.assert_not_called()
        showerror.assert_not_called()

    def test_show_batch_summary_dialog_uses_error_when_all_files_fail(self) -> None:
        failed_file = main.FailedFile(
            input_file=Path("/tmp/second.csv"),
            error=main.CsvFormatError("missing header"),
        )

        with (
            mock.patch.object(main, "Tk", return_value=FakeRoot()),
            mock.patch.object(main.messagebox, "showinfo") as showinfo,
            mock.patch.object(main.messagebox, "showwarning") as showwarning,
            mock.patch.object(main.messagebox, "showerror") as showerror,
        ):
            main.show_batch_summary_dialog([], [failed_file])

        showerror.assert_called_once()
        self.assertEqual(showerror.call_args.kwargs["title"], "Parsing failed")
        self.assertEqual(showerror.call_args.kwargs["message"], "Parsed 0 of 1 file.")
        showinfo.assert_not_called()
        showwarning.assert_not_called()

    def test_main_returns_without_parsing_when_no_files_are_selected(self) -> None:
        with (
            mock.patch.object(main, "select_input_files", return_value=[]),
            mock.patch.object(main, "parse_input_files") as parse_input_files,
            mock.patch.object(main, "show_batch_summary_dialog") as show_dialog,
            mock.patch("builtins.print") as print_message,
        ):
            main.main()

        parse_input_files.assert_not_called()
        show_dialog.assert_not_called()
        print_message.assert_called_once_with("No files selected.")

    def write_file(self, file_path: Path, content: str) -> Path:
        file_path.write_text(content, encoding="utf-8")

        return file_path


if __name__ == "__main__":
    unittest.main()

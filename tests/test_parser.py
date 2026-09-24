from __future__ import annotations

import csv
import io
import tempfile
import unittest
from collections import OrderedDict
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

import main


HEADER = [
    main.DATE_HEADER,
    main.NAME_HEADER,
    main.ACCOUNT_HEADER,
    main.CATEGORY_HEADER,
    main.AMOUNT_HEADER,
]
COLUMNS = main.TransactionColumnIndexes(date=0, name=1, amount=4)


class CsvReadingTest(unittest.TestCase):
    def test_read_csv_rows_supports_common_encodings_and_delimiters(self) -> None:
        cases = [
            ("utf-8-sig", ";", "Płatność kartą"),
            ("cp1250", "\t", "Przelew zewnętrzny"),
            ("iso-8859-2", ",", "Zwrot opłaty"),
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)

            for encoding, delimiter, transaction_name in cases:
                with self.subTest(encoding=encoding, delimiter=delimiter):
                    input_file = temp_path / f"history-{encoding}.csv"
                    csv_content = self.make_csv_content(
                        delimiter,
                        [
                            HEADER,
                            [
                                "2026-02-03",
                                transaction_name,
                                "12 345",
                                "Zakupy",
                                "-12,34 PLN",
                            ],
                        ],
                    )
                    input_file.write_text(csv_content, encoding=encoding)

                    self.assertEqual(
                        main.read_csv_rows(input_file),
                        [
                            HEADER,
                            [
                                "2026-02-03",
                                transaction_name,
                                "12 345",
                                "Zakupy",
                                "-12,34 PLN",
                            ],
                        ],
                    )

    def test_extract_transaction_rows_skips_metadata_and_cleans_rows(self) -> None:
        csv_rows = [
            ["metadata"],
            [" generated at ", " 2026-01-01 "],
            [
                f" {main.DATE_HEADER} ",
                f" {main.NAME_HEADER} ",
                f" {main.ACCOUNT_HEADER} ",
                f" {main.CATEGORY_HEADER} ",
                f" {main.AMOUNT_HEADER} ",
                "",
            ],
            [" 2026-01-02 ", " Shop   Card ", " 123 ", " Food ", " -9,99 PLN ", ""],
            ["", "", ""],
        ]

        self.assertEqual(
            main.extract_transaction_rows(csv_rows),
            [
                HEADER,
                ["2026-01-02", "Shop   Card", "123", "Food", "-9,99 PLN"],
            ],
        )

    def test_read_csv_rows_rejects_files_without_transaction_table(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            input_file = Path(temp_dir) / "invalid.csv"
            input_file.write_text("not an mbank export\n1;2;3\n", encoding="utf-8")

            with self.assertRaisesRegex(
                main.CsvFormatError,
                "Could not find the mBank transaction table header",
            ):
                main.read_csv_rows(input_file)

    def test_transaction_column_indexes_rejects_missing_required_columns(self) -> None:
        rows = [
            [
                main.DATE_HEADER,
                main.NAME_HEADER,
                main.ACCOUNT_HEADER,
                main.CATEGORY_HEADER,
            ],
            ["2026-01-02", "Shop", "123", "Food"],
        ]

        with self.assertRaisesRegex(
            main.CsvFormatError,
            "The transaction table is missing required columns",
        ) as error:
            main.transaction_column_indexes(rows)

        self.assertIn(main.AMOUNT_HEADER, str(error.exception))

    def make_csv_content(self, delimiter: str, rows: list[list[str]]) -> str:
        output = io.StringIO()
        writer = csv.writer(output, delimiter=delimiter, lineterminator="\n")
        writer.writerows(rows)

        return output.getvalue()


class TransactionTransformationTest(unittest.TestCase):
    def test_group_transactions_by_month_uses_polish_month_names(self) -> None:
        rows = [
            HEADER,
            ["2026-01-31", "January", "123", "Other", "-1,00 PLN"],
            ["2026-02-01", "February", "123", "Other", "-2,00 PLN"],
            ["2026-01-01", "January again", "123", "Other", "-3,00 PLN"],
        ]

        self.assertEqual(
            main.group_transactions_by_month(rows, COLUMNS),
            OrderedDict(
                {
                    "Styczeń 2026": [
                        ["2026-01-31", "January", "123", "Other", "-1,00 PLN"],
                        [
                            "2026-01-01",
                            "January again",
                            "123",
                            "Other",
                            "-3,00 PLN",
                        ],
                    ],
                    "Luty 2026": [
                        ["2026-02-01", "February", "123", "Other", "-2,00 PLN"],
                    ],
                }
            ),
        )

    def test_amount_parsing_splitting_summing_and_formatting(self) -> None:
        cost = ["2026-01-01", "Cost", "123", "Other", "-1 234,50 PLN"]
        refund = ["2026-01-02", "Refund", "123", "Other", "99,40 PLN"]
        zero = ["2026-01-03", "Zero", "123", "Other", "0,00 PLN"]

        self.assertEqual(main.parse_amount(cost, COLUMNS), Decimal("-1234.50"))
        self.assertEqual(main.parse_amount(refund, COLUMNS), Decimal("99.40"))
        self.assertEqual(
            main.split_transactions_by_amount([cost, refund, zero], COLUMNS),
            ([cost], [refund]),
        )
        self.assertEqual(
            main.sum_transactions([cost, refund], COLUMNS),
            Decimal("1333.90"),
        )
        self.assertEqual(main.format_amount(Decimal("-1234.5")), "1234,50")

    def test_transaction_name_cleanup_trims_at_repeated_whitespace(self) -> None:
        self.assertEqual(
            main.format_transaction_name("  Grocery store    CARD PAYMENT  "),
            "Grocery store",
        )
        self.assertEqual(
            main.format_transaction_name("Single spaces stay"),
            "Single spaces stay",
        )

    def test_output_file_for_adds_parsed_suffix_before_extension(self) -> None:
        self.assertEqual(
            main.output_file_for(Path("/tmp/history.backup.csv")),
            Path("/tmp/history.backup_parsed.xlsx"),
        )


class XlsxOutputTest(unittest.TestCase):
    def test_write_xlsx_creates_month_sheets_with_summaries_and_tables(self) -> None:
        rows = [
            HEADER,
            ["2026-04-12", "Zwrot", "12 345", "Zwroty", "50,00 PLN"],
            [
                "2026-04-11",
                "Sklep spozywczy    KARTA",
                "12 345",
                "Zakupy",
                "-309,79 PLN",
            ],
            ["2026-04-09", "Kawiarnia", "12 345", "Jedzenie", "-10,00 PLN"],
            ["2026-05-01", "Pensja", "12 345", "Przychody", "1000,00 PLN"],
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            output_file = Path(temp_dir) / "history_parsed.xlsx"

            main.write_xlsx(rows, output_file)

            workbook = load_workbook(output_file)
            self.assertEqual(workbook.sheetnames, ["Kwiecień 2026", "Maj 2026"])

            april = workbook["Kwiecień 2026"]
            self.assertEqual(april["A1"].value, "sum costs")
            self.assertEqual(april["A2"].value, "319,79")
            self.assertEqual(april["E1"].value, "sum returns")
            self.assertEqual(april["E2"].value, "50,00")
            self.assertEqual(april["A4"].value, "costs")
            self.assertEqual(april["A5"].value, "name")
            self.assertEqual(april["B5"].value, "amount")
            self.assertEqual(april["C5"].value, "date")
            self.assertEqual(april["A6"].value, "Kawiarnia")
            self.assertEqual(april["B6"].value, "10,00")
            self.assertEqual(april["C6"].value, "09.04.2026")
            self.assertEqual(april["A7"].value, "Sklep spozywczy")
            self.assertEqual(april["B7"].value, "309,79")
            self.assertEqual(april["C7"].value, "11.04.2026")
            self.assertEqual(april["E4"].value, "returns")
            self.assertEqual(april["E5"].value, "name")
            self.assertEqual(april["F5"].value, "amount")
            self.assertEqual(april["G5"].value, "date")
            self.assertEqual(april["E6"].value, "Zwrot")
            self.assertEqual(april["F6"].value, "50,00")
            self.assertEqual(april["G6"].value, "12.04.2026")

            may = workbook["Maj 2026"]
            self.assertEqual(may["A2"].value, "0,00")
            self.assertEqual(may["E2"].value, "1000,00")
            self.assertEqual(may["E6"].value, "Pensja")
            self.assertEqual(may["F6"].value, "1000,00")
            self.assertEqual(may["G6"].value, "01.05.2026")

    def test_write_xlsx_handles_header_only_transaction_table(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output_file = Path(temp_dir) / "empty_parsed.xlsx"

            main.write_xlsx([HEADER], output_file)

            workbook = load_workbook(output_file)
            self.assertEqual(workbook.sheetnames, ["Transactions"])

            sheet = workbook["Transactions"]
            self.assertEqual(sheet["A1"].value, "sum costs")
            self.assertEqual(sheet["A2"].value, "0,00")
            self.assertEqual(sheet["E1"].value, "sum returns")
            self.assertEqual(sheet["E2"].value, "0,00")
            self.assertEqual(sheet["A4"].value, "costs")
            self.assertEqual(sheet["A5"].value, "name")
            self.assertEqual(sheet["B5"].value, "amount")
            self.assertEqual(sheet["C5"].value, "date")
            self.assertEqual(sheet["E4"].value, "returns")
            self.assertEqual(sheet["E5"].value, "name")
            self.assertEqual(sheet["F5"].value, "amount")
            self.assertEqual(sheet["G5"].value, "date")
            self.assertIsNone(sheet["A6"].value)
            self.assertIsNone(sheet["E6"].value)


if __name__ == "__main__":
    unittest.main()

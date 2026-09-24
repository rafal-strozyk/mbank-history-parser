from __future__ import annotations

import csv
import re
from collections import OrderedDict
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path
from tkinter import Tk, filedialog, messagebox

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font


ENCODINGS_TO_TRY = ("utf-8-sig", "utf-8", "cp1250", "iso-8859-2")
CSV_SNIFF_SAMPLE_SIZE = 4096
REPEATED_WHITESPACE_PATTERN = re.compile(r"\s{2,}")
LAST_COLUMN_INDEX = -1
DATE_HEADER = "#Data operacji"
NAME_HEADER = "#Opis operacji"
ACCOUNT_HEADER = "#Rachunek"
CATEGORY_HEADER = "#Kategoria"
AMOUNT_HEADER = "#Kwota"
REQUIRED_TRANSACTION_HEADERS = (
    DATE_HEADER,
    NAME_HEADER,
    ACCOUNT_HEADER,
    CATEGORY_HEADER,
    AMOUNT_HEADER,
)
OUTPUT_TABLE_HEADER = ["name", "amount", "date"]
COSTS_TABLE_TITLE = "costs"
RETURNS_TABLE_TITLE = "returns"
COSTS_TABLE_START_COLUMN = 1
RETURNS_TABLE_START_COLUMN = 5
SUMMARY_TITLE_ROW = 1
SUMMARY_AMOUNT_ROW = 2
TRANSACTION_TABLE_START_ROW = 4
OUTPUT_NAME_COLUMN_OFFSET = 0
OUTPUT_AMOUNT_COLUMN_OFFSET = 1
OUTPUT_DATE_COLUMN_OFFSET = 2
COLUMN_WIDTH_PADDING = 1
POLISH_MONTH_NAMES = {
    1: "Styczeń",
    2: "Luty",
    3: "Marzec",
    4: "Kwiecień",
    5: "Maj",
    6: "Czerwiec",
    7: "Lipiec",
    8: "Sierpień",
    9: "Wrzesień",
    10: "Październik",
    11: "Listopad",
    12: "Grudzień",
}


class CsvFormatError(Exception):
    pass


@dataclass(frozen=True)
class TransactionColumnIndexes:
    date: int
    name: int
    amount: int


@dataclass(frozen=True)
class ParsedFile:
    input_file: Path
    output_file: Path
    row_count: int


@dataclass(frozen=True)
class FailedFile:
    input_file: Path
    error: Exception


def select_input_files() -> list[Path]:
    root = Tk()
    root.withdraw()
    root.update()

    selected_files = filedialog.askopenfilenames(
        title="Select mBank transaction history CSV files",
        filetypes=[
            ("CSV files", "*.csv"),
            ("All files", "*.*"),
        ],
    )

    root.destroy()

    return [Path(selected_file) for selected_file in selected_files]


def read_csv_rows(input_file: Path) -> list[list[str]]:
    last_error: UnicodeDecodeError | None = None

    for encoding in ENCODINGS_TO_TRY:
        try:
            content = input_file.read_text(encoding=encoding)
        except UnicodeDecodeError as error:
            last_error = error
            continue

        sample = content[:CSV_SNIFF_SAMPLE_SIZE]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";"

        csv_rows = list(csv.reader(content.splitlines(), dialect))
        return extract_transaction_rows(csv_rows)

    if last_error is not None:
        raise last_error

    return []


def extract_transaction_rows(csv_rows: list[list[str]]) -> list[list[str]]:
    for row_index, row in enumerate(csv_rows):
        cleaned_row = clean_row(row)

        if is_transaction_header_row(cleaned_row):
            return [
                cleaned_row,
                *[
                    transaction_row
                    for source_row in csv_rows[row_index + 1 :]
                    if (transaction_row := clean_row(source_row))
                ],
            ]

    raise CsvFormatError(
        "Could not find the mBank transaction table header.\n\n"
        f"Expected header row:\n{';'.join(REQUIRED_TRANSACTION_HEADERS)}"
    )


def clean_row(row: list[str]) -> list[str]:
    cleaned_row = [cell.strip() for cell in row]

    remove_trailing_empty_cells(cleaned_row)

    return cleaned_row


def is_transaction_header_row(row: list[str]) -> bool:
    return DATE_HEADER in row


def remove_trailing_empty_cells(row: list[str]) -> None:
    while row and row[LAST_COLUMN_INDEX] == "":
        row.pop()


def write_xlsx(rows: list[list[str]], output_file: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    column_indexes = transaction_column_indexes(rows)

    for sheet_name, transactions in group_transactions_by_month(rows, column_indexes).items():
        sheet = workbook.create_sheet(title=sheet_name)
        write_month_sheet(sheet, transactions, column_indexes)

    workbook.save(output_file)


def transaction_column_indexes(rows: list[list[str]]) -> TransactionColumnIndexes:
    if not rows:
        raise CsvFormatError("The selected file does not contain transaction rows.")

    header = rows[0]
    missing_headers = [
        required_header
        for required_header in REQUIRED_TRANSACTION_HEADERS
        if required_header not in header
    ]

    if missing_headers:
        raise CsvFormatError(
            "The transaction table is missing required columns:\n"
            + "\n".join(missing_headers)
        )

    return TransactionColumnIndexes(
        date=header.index(DATE_HEADER),
        name=header.index(NAME_HEADER),
        amount=header.index(AMOUNT_HEADER),
    )


def write_month_sheet(
    sheet,
    transactions: list[list[str]],
    column_indexes: TransactionColumnIndexes,
) -> None:
    sorted_transactions = sort_transactions_by_date(transactions, column_indexes)
    costs, returns = split_transactions_by_amount(sorted_transactions, column_indexes)
    write_summary(sheet, "sum costs", costs, column_indexes, COSTS_TABLE_START_COLUMN)
    write_summary(sheet, "sum returns", returns, column_indexes, RETURNS_TABLE_START_COLUMN)
    write_transaction_table(
        sheet,
        COSTS_TABLE_TITLE,
        costs,
        column_indexes,
        start_row=TRANSACTION_TABLE_START_ROW,
        start_column=COSTS_TABLE_START_COLUMN,
    )
    write_transaction_table(
        sheet,
        RETURNS_TABLE_TITLE,
        returns,
        column_indexes,
        start_row=TRANSACTION_TABLE_START_ROW,
        start_column=RETURNS_TABLE_START_COLUMN,
    )
    adjust_output_column_widths(sheet)


def write_summary(
    sheet,
    title: str,
    transactions: list[list[str]],
    column_indexes: TransactionColumnIndexes,
    start_column: int,
) -> None:
    sheet.cell(row=SUMMARY_TITLE_ROW, column=start_column, value=title).font = Font(
        bold=True
    )
    sheet.cell(
        row=SUMMARY_AMOUNT_ROW,
        column=start_column,
        value=format_amount(sum_transactions(transactions, column_indexes)),
    )


def sort_transactions_by_date(
    transactions: list[list[str]],
    column_indexes: TransactionColumnIndexes,
) -> list[list[str]]:
    return sorted(
        transactions,
        key=lambda transaction: date.fromisoformat(transaction[column_indexes.date]),
    )


def write_transaction_table(
    sheet,
    title: str,
    transactions: list[list[str]],
    column_indexes: TransactionColumnIndexes,
    start_row: int,
    start_column: int,
) -> int:
    sheet.cell(row=start_row, column=start_column, value=title).font = Font(
        bold=True
    )
    header_row = start_row + 1

    for column_index, header in enumerate(OUTPUT_TABLE_HEADER, start=1):
        sheet.cell(
            row=header_row,
            column=start_column + column_index - 1,
            value=header,
        ).font = Font(bold=True)

    current_row = header_row + 1

    for transaction in transactions:
        sheet.cell(
            row=current_row,
            column=start_column + OUTPUT_NAME_COLUMN_OFFSET,
            value=format_transaction_name(transaction[column_indexes.name]),
        )
        sheet.cell(
            row=current_row,
            column=start_column + OUTPUT_AMOUNT_COLUMN_OFFSET,
            value=format_amount(parse_amount(transaction, column_indexes)),
        )
        sheet.cell(
            row=current_row,
            column=start_column + OUTPUT_DATE_COLUMN_OFFSET,
            value=format_date(transaction[column_indexes.date]),
        )
        current_row += 1

    return current_row - 1


def adjust_output_column_widths(sheet) -> None:
    for column_index in output_columns_to_autofit():
        column_letter = get_column_letter(column_index)
        max_value_length = max(
            len(str(cell.value))
            for cell in sheet[column_letter]
            if cell.value is not None
        )
        sheet.column_dimensions[column_letter].width = (
            max_value_length + COLUMN_WIDTH_PADDING
        )


def output_columns_to_autofit() -> tuple[int, ...]:
    return (
        COSTS_TABLE_START_COLUMN + OUTPUT_NAME_COLUMN_OFFSET,
        COSTS_TABLE_START_COLUMN + OUTPUT_DATE_COLUMN_OFFSET,
        RETURNS_TABLE_START_COLUMN + OUTPUT_NAME_COLUMN_OFFSET,
        RETURNS_TABLE_START_COLUMN + OUTPUT_DATE_COLUMN_OFFSET,
    )


def sum_transactions(
    transactions: list[list[str]],
    column_indexes: TransactionColumnIndexes,
) -> Decimal:
    return sum(
        (abs(parse_amount(transaction, column_indexes)) for transaction in transactions),
        Decimal("0"),
    )


def split_transactions_by_amount(
    transactions: list[list[str]],
    column_indexes: TransactionColumnIndexes,
) -> tuple[list[list[str]], list[list[str]]]:
    costs = []
    returns = []

    for transaction in transactions:
        amount = parse_amount(transaction, column_indexes)

        if amount < 0:
            costs.append(transaction)
        elif amount > 0:
            returns.append(transaction)

    return costs, returns


def parse_amount(
    transaction: list[str],
    column_indexes: TransactionColumnIndexes,
) -> Decimal:
    amount = transaction[column_indexes.amount]
    normalized_amount = amount.replace("PLN", "").replace(" ", "").replace(",", ".")

    return Decimal(normalized_amount)


def format_amount(amount: Decimal) -> str:
    return f"{abs(amount):.2f}".replace(".", ",")


def format_date(transaction_date: str) -> str:
    return date.fromisoformat(transaction_date).strftime("%d.%m.%Y")


def format_transaction_name(transaction_name: str) -> str:
    stripped_transaction_name = transaction_name.strip()
    repeated_whitespace_match = REPEATED_WHITESPACE_PATTERN.search(
        stripped_transaction_name
    )

    if repeated_whitespace_match is None:
        return stripped_transaction_name

    return stripped_transaction_name[: repeated_whitespace_match.start()]


def group_transactions_by_month(
    rows: list[list[str]],
    column_indexes: TransactionColumnIndexes,
) -> OrderedDict[str, list[list[str]]]:
    if not rows:
        return OrderedDict({"Transactions": []})

    transactions = rows[1:]
    grouped_transactions: OrderedDict[str, list[list[str]]] = OrderedDict()

    for transaction in transactions:
        sheet_name = sheet_name_for_transaction(transaction, column_indexes)
        grouped_transactions.setdefault(sheet_name, []).append(transaction)

    if not grouped_transactions:
        return OrderedDict({"Transactions": []})

    return grouped_transactions


def sheet_name_for_transaction(
    transaction: list[str],
    column_indexes: TransactionColumnIndexes,
) -> str:
    transaction_date = date.fromisoformat(transaction[column_indexes.date])
    return f"{POLISH_MONTH_NAMES[transaction_date.month]} {transaction_date.year}"


def output_file_for(input_file: Path) -> Path:
    return input_file.with_name(f"{input_file.stem}_parsed.xlsx")


def parse_input_file(input_file: Path) -> ParsedFile:
    rows = read_csv_rows(input_file)
    output_file = output_file_for(input_file)
    write_xlsx(rows, output_file)

    return ParsedFile(
        input_file=input_file,
        output_file=output_file,
        row_count=len(rows),
    )


def parse_input_files(
    input_files: list[Path],
) -> tuple[list[ParsedFile], list[FailedFile]]:
    parsed_files = []
    failed_files = []

    for input_file in input_files:
        try:
            parsed_files.append(parse_input_file(input_file))
        except Exception as error:
            failed_files.append(FailedFile(input_file=input_file, error=error))

    return parsed_files, failed_files


def show_batch_summary_dialog(
    parsed_files: list[ParsedFile],
    failed_files: list[FailedFile],
) -> None:
    root = Tk()
    root.withdraw()

    selected_file_count = len(parsed_files) + len(failed_files)
    summary = (
        f"Parsed {len(parsed_files)} of {selected_file_count} "
        f"{pluralize('file', selected_file_count)}."
    )
    detail = batch_summary_detail(parsed_files, failed_files)

    if not failed_files:
        messagebox.showinfo(
            title="Parsing complete",
            message=summary,
            detail=detail,
            parent=root,
        )
    elif parsed_files:
        messagebox.showwarning(
            title="Parsing completed with errors",
            message=summary,
            detail=detail,
            parent=root,
        )
    else:
        messagebox.showerror(
            title="Parsing failed",
            message=summary,
            detail=detail,
            parent=root,
        )

    root.destroy()


def batch_summary_detail(
    parsed_files: list[ParsedFile],
    failed_files: list[FailedFile],
) -> str:
    detail_sections = []

    if parsed_files:
        detail_sections.append(
            "Created outputs:\n"
            + "\n".join(
                f"- {parsed_file.output_file}" for parsed_file in parsed_files
            )
        )

    if failed_files:
        detail_sections.append(
            "Failed files:\n"
            + "\n".join(
                f"- {failed_file.input_file}: {format_error(failed_file.error)}"
                for failed_file in failed_files
            )
        )

    return "\n\n".join(detail_sections)


def format_error(error: Exception) -> str:
    error_message = str(error).strip()

    if error_message:
        return error_message

    return error.__class__.__name__


def pluralize(word: str, count: int) -> str:
    if count == 1:
        return word

    return f"{word}s"


def main() -> None:
    input_files = select_input_files()

    if not input_files:
        print("No files selected.")
        return

    parsed_files, failed_files = parse_input_files(input_files)
    show_batch_summary_dialog(parsed_files, failed_files)

    print(f"Selected {len(input_files)} {pluralize('file', len(input_files))}.")

    for parsed_file in parsed_files:
        print(
            f"Wrote {parsed_file.row_count} rows from "
            f"{parsed_file.input_file} to: {parsed_file.output_file}"
        )

    for failed_file in failed_files:
        print(f"Failed to parse {failed_file.input_file}: {failed_file.error}")


if __name__ == "__main__":
    main()
